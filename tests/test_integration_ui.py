"""真机端到端 UI 测试：**真的用浏览器渲染一遍**，而不是只验 HTML 合法。

为什么必须有这一层
------------------
观察台和任务窗格是这个项目的主可见面，但我一度只验到"HTML 合法、JS 语法正确、
接口有数据"。这三条加起来**仍然不等于页面对**。第一次真截图就看到三处问题：

1. 「可用工具」永远是空的 —— `loadTools()` 忘了带令牌，401 被 catch 悄悄吞掉；
2. 窄窗格里「急停」按钮"不见了" —— 一度误判成布局 bug，实际是**我的截图方法**在裁剪：
   headless 有约 518px 的最小窗口宽度，`--window-size=380` 只是把宽布局裁成 380；
3. 修问题的过程中**我自己引入了一个回归** —— `addRow(id, dotClass, text, sub)` 参数传错，
   日志行的文字全没了。也是靠再看一眼截图才发现的。

所以这个文件走 CDP（`tests/cdp.py`）：**确定性等待 + 真实视口**，不看运气。

    .venv\\Scripts\\python.exe -m pytest -m com -v tests/test_integration_ui.py
"""

from __future__ import annotations

import sys
import os
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from pptd.config import read_runtime, settings as cfg  # noqa: E402

cdp = pytest.importorskip("cdp", reason="UI 测试需要 tests/cdp.py")

if cdp.find_browser() is None:
    pytest.skip("本机没有 Edge/Chrome，跳过 UI 渲染测试", allow_module_level=True)

pytestmark = pytest.mark.com


def _origin() -> str:
    return f"http://{cfg.host}:{cfg.port}"


@pytest.fixture(scope="module")
def daemon():
    from pptd.daemon import ensure, is_alive, stop

    already = is_alive(cfg)
    info = ensure(cfg, timeout=40)
    yield info
    if not already:
        stop(cfg, force=True)


@pytest.fixture(scope="module")
def prepared_deck(daemon: dict, tmp_path_factory):
    """确保有演示打开、且跑过几步操作——否则页面没什么可渲染的。"""
    import httpx
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    token = str(read_runtime(cfg)["token"])
    sample = tmp_path_factory.mktemp("ui-demo") / "ppt-agent-ui-demo.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for index, title in enumerate(("真实 PowerPoint", "读取与编辑", "操作全程可见", "同一份工具契约", "让每一步都看得见"), 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("141B2B")
        for text, y, size, color in (
            ("ppt-agent  /  LIVE POWERPOINT", 0.8, 18, "4DD9E8"),
            (title, 2.1, 42, "FFFFFF"),
            ("AI agent → COM → PowerPoint\n读取 · 编辑 · 回读验证 · 实时取景", 3.5, 23, "BCC7DA"),
            (f"演示示例  /  {index:02d}", 6.4, 15, "FF9A55"),
        ):
            box = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(11.5), Inches(1.2))
            box.text = text
            for paragraph in box.text_frame.paragraphs:
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor.from_string(color)
    deck.save(sample)
    with httpx.Client(timeout=90) as client:
        headers = {"X-PPT-Token": token}

        def call(name: str, args: dict) -> dict:
            response = client.post(f"{_origin()}/call", json={"name": name, "args": args}, headers=headers)
            response.raise_for_status()
            result = response.json()
            assert result.get("ok"), result
            return result

        call("ppt_open", {"path": str(sample)})
        call("ppt_goto", {"slide": 5})
        # 清出旧事件，使公开截图仅包含自建示例演示的操作。
        for _ in range(12):
            call("ppt_status", {})
            call("ppt_read_slide", {"slide": 5})
        call("ppt_set_text", {"slide": 5, "shape": "2", "text": "让每一步都看得见"})
        call("ppt_goto", {"slide": 5})
        try:
            yield token
        finally:
            call("ppt_close", {"deck": str(sample), "save": False})


@pytest.fixture(scope="module")
def browser():
    with cdp.Browser() as instance:
        yield instance


# --------------------------------------------------------------------------
# 观察台
# --------------------------------------------------------------------------


@pytest.fixture(scope="module")
def observer(browser, prepared_deck) -> "cdp.Browser":
    browser.goto(f"{_origin()}/")
    # 确定性等待：工具列表填好了才算页面就绪（这正是 --dump-dom 给不了的）
    browser.wait_for("document.querySelectorAll('#tools li').length > 0", timeout=25)
    return browser


def test_observer_tool_list_is_populated(observer) -> None:
    """回归：`loadTools()` 一度忘了带令牌，工具列表永远是空的。

    HTML 合法、JS 语法正确、接口也通——三条都对，页面就是空的。
    """
    names = observer.evaluate("Array.from(document.querySelectorAll('#tools code')).map(e => e.textContent)")
    assert len(names) >= 20, f"工具列表只有 {len(names)} 项"
    for must in ("ppt_decks", "ppt_read_slide", "ppt_shot", "ppt_open", "ppt_undo"):
        assert must in names, f"工具列表缺 {must}"
    assert observer.evaluate("document.querySelector('#tools').textContent.includes('读取工具列表失败')") is False


def test_observer_status_panel_is_filled(observer) -> None:
    text = observer.evaluate("document.getElementById('state').textContent")
    assert "PowerPoint" in text and "16.0" in text
    assert "ppt-agent-ui-demo.pptx" in text, "状态面板没显示当前演示"
    assert "接管已运行的实例" in text or "由 pptd 启动" in text
    assert "\\" not in text.split("当前演示")[-1][:60], "当前演示不该铺一整条路径"


def test_observer_feed_shows_chinese_human_summaries(observer) -> None:
    rows = observer.evaluate("Array.from(document.querySelectorAll('#feed li')).map(li => li.textContent)")
    assert rows, "操作流是空的"
    joined = " ".join(rows)
    assert any(word in joined for word in ("已切到第", "已打开《", "正常，打开着")), joined[:400]


def test_observer_log_rows_keep_their_text(observer) -> None:
    """回归：`addRow(id, dotClass, text, sub)` 参数传错，日志行的文字全没了、只剩时间和圆点。

    刻意**不**断言"某条具体日志还在"——事件窗口只有最近 40 条，前面别的测试跑过一堆操作后
    启动日志会被挤出去（第一版就这么写，单跑通过、整套跑失败）。
    这里查的是结构：**每一行都必须有文字**，与"当前有哪些事件"无关。
    """
    texts = observer.evaluate(
        "Array.from(document.querySelectorAll('#feed li div.m > div:first-child'))"
        ".map(e => e.textContent.trim())"
    )
    assert texts, "操作流是空的"
    empty = [i for i, text in enumerate(texts) if not text]
    assert not empty, f"第 {empty} 行没有文字（addRow 参数传错会这样）"
    assert any(("已切到第" in t) or ("已打开" in t) or ("正常，打开着" in t) or ("已截取" in t) for t in texts), (
        f"操作流里没有中文人话摘要：{texts[:5]}"
    )


def test_observer_live_view_actually_renders_the_slide(observer) -> None:
    """实时取景真的出图了——这是"用户能看见 agent 操作 PowerPoint"的硬证据。

    必须**等**出图，不能读一眼就断言：``/live`` 要走一次 PowerPoint 导出，通常 1~3 秒。
    第一版直接读 ``stage-cap``，单跑能过、整套跑就挂——前面别的测试改动了状态，
    取景更慢。**同一条测试单跑过、整套挂，就是缺等待。**
    """
    observer.wait_for(
        "(() => { const i = document.getElementById('stage-img'); return i && i.naturalWidth > 100; })()",
        timeout=30,
    )
    caption = observer.evaluate("document.getElementById('stage-cap').textContent")
    assert "实时" in caption, f"大屏没进入实时模式：{caption}"
    natural = observer.evaluate(
        "(() => { const i = document.getElementById('stage-img');"
        " return { w: i.naturalWidth, h: i.naturalHeight, shown: i.style.display }; })()"
    )
    assert natural["shown"] == "block", natural
    assert natural["w"] > 100 and natural["h"] > 100, f"舞台图没有真正解码：{natural}"


def test_observer_screenshot_is_not_blank(observer, tmp_path: Path) -> None:
    """整页截图要有真实内容——纯色页面说明渲染挂了。"""
    from PIL import Image

    observer.wait_for("document.getElementById('stage-img').naturalWidth > 100", timeout=30)
    path = observer.screenshot(tmp_path / "observer.png", width=1440, height=900, full_page=False)
    if os.environ.get("PPT_AGENT_UPDATE_SCREENSHOTS") == "1":
        import shutil
        shutil.copyfile(path, ROOT / "docs/screenshots/observer.png")
    with Image.open(path) as image:
        assert image.width >= 1400
        colors = image.convert("RGB").getcolors(maxcolors=2_000_000)
        assert colors and len(colors) > 200, f"截图颜色数只有 {len(colors or [])}，页面可能是空的"


# --------------------------------------------------------------------------
# 任务窗格：窄版式
# --------------------------------------------------------------------------


@pytest.fixture()
def pane(browser, prepared_deck) -> "cdp.Browser":
    browser.set_viewport(380, 760)
    browser.goto(f"{_origin()}/addin/")
    browser.wait_for("document.querySelectorAll('#feed li').length > 0", timeout=25)
    return browser


@pytest.mark.parametrize("width", [360, 380, 420])
def test_pane_footer_buttons_stay_inside_viewport(browser, prepared_deck, width: int) -> None:
    """回归：窄窗格里「急停」被挤出可视区。

    用 CDP 的视口覆写拿到**真实**的窄视口——headless 的 `--window-size` 有约 518px 下限，
    直接开 380 只会拿到被裁剪的宽布局（我为此误判过一轮）。
    """
    browser.set_viewport(width, 760)
    browser.goto(f"{_origin()}/addin/")
    layout = browser.evaluate(
        """(() => {
          const stop = document.getElementById('btnStop').getBoundingClientRect();
          const observer = document.getElementById('btnObserver').getBoundingClientRect();
          const footer = document.querySelector('footer').getBoundingClientRect();
          return {
            innerW: window.innerWidth,
            stopRight: Math.round(stop.right), stopBottom: Math.round(stop.bottom),
            observerRight: Math.round(observer.right),
            footerBottom: Math.round(footer.bottom),
            innerH: window.innerHeight,
          };
        })()"""
    )
    assert layout["innerW"] == width, f"视口没被真正设成 {width}：{layout}"
    assert layout["stopRight"] <= width + 1, f"{width}px 下「急停」被挤出右边：{layout}"
    assert layout["observerRight"] <= width + 1, f"{width}px 下「观察台」被挤出右边：{layout}"
    assert layout["stopBottom"] <= layout["innerH"] + 1, f"{width}px 下底栏落到视口之下：{layout}"


def test_pane_does_not_overflow_horizontally(pane) -> None:
    """长文本（URL、路径）必须折行，而不是把窄面板撑出横向滚动。"""
    overflow = pane.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth")
    assert overflow <= 1, f"页面横向溢出 {overflow}px"


def test_pane_status_line_is_readable(pane) -> None:
    status = pane.evaluate("document.getElementById('statusText').textContent")
    sub = pane.evaluate("document.getElementById('sub').textContent")
    assert "PowerPoint" in status or "未连接" in status, status
    assert "ppt-agent-ui-demo.pptx" in sub or "第" in sub, sub


def test_pane_live_view_renders(pane) -> None:
    pane.wait_for(
        "(() => { const i = document.getElementById('stageImg'); return i && i.naturalWidth > 100; })()",
        timeout=25,
    )
    size = pane.evaluate(
        "(() => { const i = document.getElementById('stageImg');"
        " return { w: i.naturalWidth, h: i.naturalHeight }; })()"
    )
    assert size["w"] > 100 and size["h"] > 100, size
    if os.environ.get("PPT_AGENT_UPDATE_SCREENSHOTS") == "1":
        pane.screenshot(ROOT / "docs/screenshots/taskpane.png", width=380, height=760, full_page=False)


def test_pane_uses_same_origin_for_observer_link(browser, prepared_deck) -> None:
    """HTTPS 页面跳 HTTP、或跨源 fetch，都会被混合内容规则拦掉，所以必须全部同源。

    注意不能直接全文搜 "http://127.0.0.1"——日志正文里**合法地**提到了守护进程地址，
    blob URL 的 origin 也长这样。要看的是**页面自己的脚本**里有没有绝对地址。
    """
    browser.set_viewport(380, 760)
    browser.goto(f"{_origin()}/addin/")
    script = browser.evaluate("document.getElementById('paneScript').textContent")
    assert "OBSERVER = '/?token='" in script, "观察台链接应当同源"
    assert "http://" not in script and "https://" not in script, (
        "脚本里出现绝对地址，跨源/混合内容会被拦掉：" + script[:200]
    )
    # 相对路径的取数端点必须在
    assert "'/feed.txt" in script or '"/feed.txt' in script
    assert "'/live" in script or '"/live' in script
